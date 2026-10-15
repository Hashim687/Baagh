"""Seeded-demo workflow tests: GST maths through real orders, Excel download, refresh-safe data, OTP rules."""
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

os.environ["BAAGH_ENV"] = "development"
os.environ["BAAGH_SEED_DEMO"] = "true"
for _k in ("TWILIO_ACCOUNT_SID", "TWILIO_AUTH_TOKEN", "TWILIO_FROM_PHONE", "TWILIO_VERIFY_SERVICE_SID"):
    os.environ.pop(_k, None)

import openpyxl  # noqa: E402
import main  # noqa: E402
import sms  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from test_core_workflows import register_member  # noqa: E402


class BaaghWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory(prefix="baagh-tests-")
        main.DB_PATH = Path(cls.temp_dir.name) / "baagh-test.sqlite3"
        main.SEED_DEMO = True
        main.OTP_RESEND_SECONDS = 0
        cls.client = TestClient(main.app)
        cls.client.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        cls.temp_dir.cleanup()

    def login(self, identifier, password):
        r = self.client.post("/api/auth/login", json={"phone": identifier, "password": password})
        self.assertEqual(r.status_code, 200, r.text)
        return {"Authorization": "Bearer " + r.json()["access_token"]}, r.json()["user"]

    def trade(self, seller_h, buyer_h, crop, price, qty):
        listing = next(x for x in self.client.get("/api/listings").json() if x["crop"] == crop)
        offer = self.client.post(f"/api/listings/{listing['id']}/offers", headers=buyer_h, json={"offered_price": price, "quantity_kg": qty})
        self.assertEqual(offer.status_code, 201, offer.text)
        order_id = self.client.post(f"/api/offers/{offer.json()['id']}/decision?decision=accept", headers=seller_h).json()["order_id"]
        self.client.post(f"/api/orders/{order_id}/dispatch", headers=seller_h)
        self.client.post(f"/api/orders/{order_id}/deliver", headers=seller_h)
        return order_id

    def ensure_trades(self):
        """Create the shared set of delivered trades once, whichever test runs first."""
        cls = type(self)
        if getattr(cls, "_traded", None):
            return cls._traded
        buyer_h, _ = self.login("9800000004", "Demo1234!")                # registered Delhi buyer
        reg_seller_h, _ = self.login("9800000003", "Demo1234!")           # GSTIN, Kupwara
        unreg_seller_h, _ = self.login("9800000002", "Demo1234!")         # no GSTIN, Baramulla
        saffron_seller_h, _ = self.login("9800000007", "Demo1234!")       # GSTIN, Pulwama
        self.trade(reg_seller_h, buyer_h, "Almond", 640, 100)             # 64,000 taxable, IGST 3,200
        self.trade(unreg_seller_h, buyer_h, "Walnut", 335, 100)           # unregistered seller => no GST
        self.trade(unreg_seller_h, buyer_h, "Apple", 95, 200)             # nil rated
        self.trade(saffron_seller_h, buyer_h, "Saffron", 295000, 1)       # 295,000 taxable, IGST 14,750
        cls._traded = (buyer_h, reg_seller_h, unreg_seller_h, saffron_seller_h)
        return cls._traded

    def test_gst_split_by_state_and_registration(self):
        buyer_h, reg_seller_h, unreg_seller_h, saffron_seller_h = self.ensure_trades()
        s = self.client.get("/api/my/gst-summary", headers=reg_seller_h).json()["summary"]
        self.assertEqual((s["taxable"], s["igst"], s["cgst"], s["sgst"]), (64000.0, 3200.0, 0.0, 0.0))
        s = self.client.get("/api/my/gst-summary", headers=saffron_seller_h).json()["summary"]
        self.assertEqual((s["taxable"], s["igst"]), (295000.0, 14750.0))
        s = self.client.get("/api/my/gst-summary", headers=unreg_seller_h).json()["summary"]
        self.assertEqual((s["total_tax"], s["taxable"]), (0, 33500.0 + 19000.0))
        b = self.client.get("/api/my/gst-summary", headers=buyer_h).json()["summary"]
        self.assertEqual(b["igst"], 3200.0 + 14750.0)
        self.assertEqual(b["invoices"], 4)

    def test_excel_report_is_tally_ready_and_balanced(self):
        buyer_h, seller_h, _, _ = self.ensure_trades()
        r = self.client.get("/api/my/gst-report.xlsx", headers=seller_h)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        for name in ("Summary", "Tally Voucher Register", "Tally Ledger Entries", "Tally Ledger Masters", "GSTR-1 B2B", "HSN Summary", "Nil-Exempt Summary"):
            self.assertIn(name, wb.sheetnames)
        balance = {}
        for row in wb["Tally Ledger Entries"].iter_rows(min_row=2, values_only=True):
            if row[2] and row[0] != "TOTAL" and isinstance(row[4], (int, float)):
                balance[row[2]] = balance.get(row[2], 0) + (row[4] if row[5] == "Dr" else -row[4])
        self.assertTrue(balance)
        self.assertTrue(all(abs(v) < 0.005 for v in balance.values()), "every voucher must balance Dr = Cr")
        # Buyers get a purchase register instead of GSTR-1 sheets.
        wb2 = openpyxl.load_workbook(io.BytesIO(self.client.get("/api/my/gst-report.xlsx", headers=buyer_h).content))
        self.assertNotIn("GSTR-1 B2B", wb2.sheetnames)
        self.assertEqual(wb2["Tally Voucher Register"]["B2"].value, "Purchase")
        self.assertEqual(self.client.get("/api/my/gst-report.xlsx?fy=2026", headers=seller_h).status_code, 422)
        self.assertEqual(self.client.get("/api/my/gst-report.xlsx").status_code, 401)

    def test_otp_rules(self):
        c = self.client
        base = {"name": "OTP Rules Co", "phone": "9811111111", "country": "IN", "password": "Buyer!Password-2026-QA",
                "role": "buyer", "state": "Punjab", "district": "Ludhiana", "pincode": "141001"}
        # Cooldown between sends
        main.OTP_RESEND_SECONDS = 60
        try:
            first = c.post("/api/auth/register", json=base)
            self.assertEqual(first.status_code, 202, first.text)
            self.assertEqual(c.post("/api/auth/resend-otp", json={"phone": base["phone"]}).status_code, 429)
        finally:
            main.OTP_RESEND_SECONDS = 0
        otp = first.json()["dev_otp"]
        # Five wrong guesses lock the code; even the right one is refused until a resend.
        wrong = "000000" if otp != "000000" else "111111"
        for _ in range(5):
            self.assertEqual(c.post("/api/auth/verify-otp", json={"phone": base["phone"], "otp": wrong}).status_code, 400)
        self.assertEqual(c.post("/api/auth/verify-otp", json={"phone": base["phone"], "otp": otp}).status_code, 429)
        again = c.post("/api/auth/resend-otp", json={"phone": base["phone"]})
        self.assertEqual(again.status_code, 200, again.text)
        self.assertEqual(c.post("/api/auth/verify-otp", json={"phone": base["phone"], "otp": again.json()["dev_otp"]}).status_code, 201)
        # Expired code
        exp = {**base, "name": "Expired Co", "phone": "9822222222"}
        code = c.post("/api/auth/register", json=exp).json()["dev_otp"]
        with main.db_conn() as conn:
            conn.execute("UPDATE pending_registrations SET expires_at=0 WHERE phone='+919822222222'")
        self.assertEqual(c.post("/api/auth/verify-otp", json={"phone": exp["phone"], "otp": code}).status_code, 400)

    def test_registration_validation(self):
        c = self.client
        good = {"name": "Validation Co", "phone": "9833333333", "country": "IN", "password": "Buyer!Password-2026-QA",
                "role": "buyer", "state": "Delhi", "district": "North", "pincode": "110033"}
        cases = {
            "short indian phone": {"phone": "98333"},
            "landline-style 5xxxx": {"phone": "5833333333"},
            "pincode of another district": {"pincode": "190001"},
            "unknown state": {"state": "Atlantis"},
            "district not in state": {"district": "Ludhiana"},
            "bad pincode format": {"pincode": "0110033"},
            "bad GSTIN": {"gstin": "07AAACB1234C1ZZ"},
            "weak password": {"password": "short"},
            "UAE number too short": {"country": "AE", "phone": "12345", "district": "Dubai"},
            "seller outside J&K": {"role": "seller"},
            "seller abroad": {"role": "seller", "country": "AE", "phone": "501234567", "district": "Dubai"},
        }
        for label, change in cases.items():
            r = c.post("/api/auth/register", json={**good, **change})
            self.assertIn(r.status_code, (422,), f"{label}: {r.status_code} {r.text}")
        # GSTIN state must match the chosen state (Punjab GSTIN used with a Delhi address).
        import gst
        g = "03AAACB1234C1Z" + gst.gstin_checksum_char("03AAACB1234C1Z")
        self.assertEqual(c.post("/api/auth/register", json={**good, "gstin": g}).status_code, 422)
        g7 = "07AAACB1234C1Z" + gst.gstin_checksum_char("07AAACB1234C1Z")
        self.assertEqual(c.post("/api/auth/register", json={**good, "gstin": g7}).status_code, 202)
        # A buyer from the UAE with a valid number is accepted.
        uae = register_member(c, name="Dubai Dry Fruits LLC", phone="501234567", country="AE", district="Dubai", state="", pincode="")
        self.assertEqual(uae.status_code, 201, uae.text)
        self.assertEqual(uae.json()["user"]["country"], "AE")
        self.assertEqual(uae.json()["user"]["phone"], "+971501234567")

    def test_production_without_sms_refuses_signup(self):
        with mock.patch.object(main, "BAAGH_ENV", "production"), mock.patch.dict(os.environ, {"BAAGH_ENV": "production"}):
            r = self.client.post("/api/auth/register", json={"name": "Prod Co", "phone": "9844444444", "country": "IN", "password": "Buyer!Password-2026-QA",
                                                              "role": "buyer", "state": "Delhi", "district": "North", "pincode": "110033"})
            self.assertEqual(r.status_code, 503)
            self.assertEqual(self.client.get("/api/auth/config").json()["sms"], "unavailable")

    def test_twilio_verify_path_is_used_when_configured(self):
        env = {"TWILIO_ACCOUNT_SID": "ACtest", "TWILIO_AUTH_TOKEN": "tok", "TWILIO_VERIFY_SERVICE_SID": "VAtest"}
        calls = []

        def fake_post(url, fields):
            calls.append((url, dict(fields)))
            if url.endswith("/VerificationCheck"):
                return 200, {"status": "approved" if fields["Code"] == "123456" else "pending"}
            return 201, {"status": "pending"}

        body = {"name": "Verify Svc Co", "phone": "9855555555", "country": "IN", "password": "Buyer!Password-2026-QA",
                "role": "buyer", "state": "Delhi", "district": "North", "pincode": "110033"}
        with mock.patch.dict(os.environ, env), mock.patch.object(sms, "_post", fake_post):
            r = self.client.post("/api/auth/register", json=body)
            self.assertEqual(r.status_code, 202, r.text)
            self.assertNotIn("dev_otp", r.json(), "real SMS mode must never expose the code")
            self.assertEqual(r.json()["delivery"], "sms")
            self.assertEqual(self.client.post("/api/auth/verify-otp", json={"phone": body["phone"], "otp": "654321"}).status_code, 400)
            ok = self.client.post("/api/auth/verify-otp", json={"phone": body["phone"], "otp": "123456"})
            self.assertEqual(ok.status_code, 201, ok.text)
        self.assertTrue(calls[0][0].endswith("/Verifications"))
        self.assertEqual(calls[0][1]["To"], "+919855555555")

    def test_catalog_and_geo_reference_data(self):
        cat = self.client.get("/api/catalog").json()
        crops = {c["crop"]: c for c in cat["crops"]}
        self.assertEqual(set(crops), {"Apple", "Walnut", "Almond", "Saffron"})
        self.assertEqual(crops["Apple"]["forms"][0]["gst_rate"], 0)
        self.assertTrue(all(f["gst_rate"] == 5 for f in crops["Walnut"]["forms"] + crops["Almond"]["forms"] + crops["Saffron"]["forms"]))
        self.assertGreaterEqual(len(crops["Apple"]["varieties"]), 15)
        kashmir = self.client.get("/api/geo/kashmir").json()
        names = {d["district"] for d in kashmir["districts"] if d["division"] == "Kashmir"}
        self.assertEqual(names, {"Anantnag", "Bandipora", "Baramulla", "Budgam", "Ganderbal", "Kulgam", "Kupwara", "Pulwama", "Shopian", "Srinagar"})
        self.assertEqual(len(self.client.get("/api/geo/states").json()), 36)
        self.assertGreater(len(self.client.get("/api/geo/countries").json()), 200)
        self.assertEqual(self.client.get("/api/geo/pincode/110033").json()[0]["state"], "Delhi")
        self.assertEqual(self.client.get("/api/geo/pincode/999999").status_code, 404)
        india = next(c for c in self.client.get("/api/geo/countries").json() if c["iso"] == "IN")
        self.assertEqual((india["dial"], india["lengths"]), ("+91", [10]))

    def test_listing_validation(self):
        h, _ = self.login("9800000007", "Demo1234!")
        ok = {"crop": "Saffron", "product_form": "Powder", "variety": "Guchhi - Grade III", "grade": "C", "quantity_kg": 0.25, "price_per_kg": 150000, "district": "Pulwama"}
        r = self.client.post("/api/listings", headers=h, json=ok)
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual((r.json()["hsn_code"], r.json()["gst_rate"], r.json()["grade_label"]), ("0910", 5, "Guchhi / Zarda (Grade III-IV)"))
        for bad in ({"crop": "Mango"}, {"product_form": "Frozen"}, {"district": "Mumbai"}):
            self.assertEqual(self.client.post("/api/listings", headers=h, json={**ok, **bad}).status_code, 422, bad)

    def test_buyer_directory_access(self):
        seller_h, _ = self.login("9800000002", "Demo1234!")
        buyer_h, _ = self.login("9800000004", "Demo1234!")
        rows = self.client.get("/api/buyers?state=Delhi", headers=seller_h).json()
        self.assertTrue(any(b["name"] == "Azad Traders" and b["gst_registered"] for b in rows))
        self.assertNotIn("phone", rows[0])
        self.assertEqual(self.client.get("/api/buyers", headers=buyer_h).status_code, 403)
        self.assertEqual(self.client.get("/api/buyers").status_code, 401)


if __name__ == "__main__":
    unittest.main(verbosity=2)
